import os
import openpyxl

SAMPLE_DIR = os.path.join(os.path.dirname(__file__), "sample_sheets")
os.makedirs(SAMPLE_DIR, exist_ok=True)

def create_ecommerce_parcels():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Parcels"
    
    headers = ["Shipment_ID", "Origin_ZIP", "Destination_ZIP", "Weight_Lbs", "Length_In", "Width_In", "Height_In", "Service_Level"]
    ws.append(headers)
    
    rows = [
        ["PKG-1001", "94103", "60601", 14.5, 12, 10, 6, "Standard Ground"],
        ["PKG-1002", "94103", "10001", 22.0, 16, 12, 8, "Priority 2-Day"],
        ["PKG-1003", "94103", "30301", 8.2, 10, 8, 4, "Ground Economy"],
        ["PKG-1004", "94103", "98101", 35.0, 20, 14, 10, "Standard Ground"],
        ["PKG-1005", "94103", "75001", 18.5, 14, 12, 8, "Priority Overnight"],
        ["PKG-1006", "94103", "02108", 5.0, 8, 6, 4, "Standard Ground"],
        ["PKG-1007", "94103", "80202", 48.0, 24, 18, 12, "Ground Heavy"]
    ]
    for r in rows:
        ws.append(r)
        
    path = os.path.join(SAMPLE_DIR, "ecommerce_parcels.xlsx")
    wb.save(path)
    print("Created:", path)

def create_wholesale_pallets():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LTL_Pallets"
    
    headers = ["Pallet_Ref", "Pickup_ZIP", "Delivery_ZIP", "Total_Weight_Lbs", "Pallet_Count", "Class", "Liftgate_Required"]
    ws.append(headers)
    
    rows = [
        ["PLT-501", "94103", "48201", 620.0, 1, 70, "Yes"],
        ["PLT-502", "94103", "37201", 1240.0, 2, 85, "No"],
        ["PLT-503", "94103", "90001", 450.0, 1, 60, "No"],
        ["PLT-504", "94103", "64101", 1850.0, 3, 92.5, "Yes"],
        ["PLT-505", "94103", "78701", 980.0, 2, 77.5, "No"]
    ]
    for r in rows:
        ws.append(r)
        
    path = os.path.join(SAMPLE_DIR, "wholesale_pallets_ltl.xlsx")
    wb.save(path)
    print("Created:", path)

def create_medical_freight():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Priority_Cargo"
    
    headers = ["Cargo_ID", "Origin_Postal", "Dest_Postal", "Weight_Lbs", "Temp_Controlled", "Priority_Rating"]
    ws.append(headers)
    
    rows = [
        ["MED-801", "94103", "19104", 42.0, "Yes", "Critical Next-Flight"],
        ["MED-802", "94103", "55401", 18.5, "No", "Express Priority"],
        ["MED-803", "94103", "33101", 65.0, "Yes", "Critical Next-Flight"],
        ["MED-804", "94103", "20001", 12.0, "No", "Standard Express"]
    ]
    for r in rows:
        ws.append(r)
        
    path = os.path.join(SAMPLE_DIR, "medical_priority_freight.xlsx")
    wb.save(path)
    print("Created:", path)

if __name__ == "__main__":
    create_ecommerce_parcels()
    create_wholesale_pallets()
    create_medical_freight()
