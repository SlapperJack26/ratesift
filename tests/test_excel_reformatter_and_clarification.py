import unittest
import io
import openpyxl
from fastapi.testclient import TestClient

from main import app
from services.auth_service import create_session
from services.import_day_and_ross_guide import import_day_and_ross_tariff

class TestExcelReformatterAndClarification(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)
        cls.user_id = "usr_alex_rivers"
        cls.token = create_session(cls.user_id)
        cls.client.cookies.set("shipflow_session", cls.token)

        # Import confirmed Day & Ross Guide Tariff for testing quotes
        cls.guide_sheet_id = import_day_and_ross_tariff(user_id=cls.user_id, auto_confirm=True)

    def _create_sample_excel(self, rows_data, headers=None):
        """Helper to create in-memory Excel files for testing."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Shipments"

        if headers:
            ws.append(headers)
        for row in rows_data:
            ws.append(row)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream.getvalue()

    def test_01_rule1_and_rule25_zero_assumption_and_missing_details(self):
        """Rule 1 & 25: Never guess missing parameters; reject unconfirmed blanks."""
        # Row 1: Valid
        # Row 2: Missing weight
        # Row 3: Missing destination
        excel_bytes = self._create_sample_excel(
            headers=["Origin", "Destination", "Weight (lbs)", "Dims"],
            rows_data=[
                ["CALGARY, AB", "LINDSAY, ON", 1450, "48x40x50"],
                ["TORONTO, ON", "MONTREAL, QC", None, "20x20x20"],
                ["EDMONTON, AB", None, 500, "10x10x10"]
            ]
        )

        resp = self.client.post(
            "/api/quotes/analyze-excel",
            files={"file": ("customer_orders.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()

        self.assertIn("analysis_token", data)
        self.assertEqual(data["total_rows"], 3)
        self.assertTrue(data["has_blockers"])

        # Row 2 must have missing weight issue
        row_2 = next(r for r in data["preview_rows"] if r["row_num"] == 3)
        self.assertFalse(row_2["is_complete"])
        self.assertTrue(row_2["needs_review"])
        self.assertIsNone(row_2["weight_lbs"])

        # Row 3 must have missing destination issue
        row_3 = next(r for r in data["preview_rows"] if r["row_num"] == 4)
        self.assertFalse(row_3["is_complete"])
        self.assertTrue(row_3["needs_review"])
        self.assertEqual(row_3["dest_normalized"], "")

        # Rule 25 Enforcement: Attempting to confirm without clarifying missing data MUST fail with 400
        token = data["analysis_token"]
        confirm_resp_fail = self.client.post("/api/quotes/confirm-reformatted", json={
            "analysis_token": token,
            "row_corrections": {}  # No corrections provided
        })
        self.assertEqual(confirm_resp_fail.status_code, 400)
        self.assertIn("Rule 25 Requirement", confirm_resp_fail.json()["detail"])

        # Supplying the missing fields allows confirmation to succeed
        confirm_resp_success = self.client.post("/api/quotes/confirm-reformatted", json={
            "analysis_token": token,
            "row_corrections": {
                "3": {"weight_lbs": 650.0},
                "4": {"dest_normalized": "WINNIPEG, MB"}
            }
        })
        self.assertEqual(confirm_resp_success.status_code, 200)
        conf_data = confirm_resp_success.json()
        self.assertEqual(conf_data["status"], "success")
        self.assertEqual(conf_data["total_shipments"], 3)

    def test_02_rule3_source_cell_coordinates_traceability(self):
        """Rule 3: Ensure genuine cell coordinates (e.g. Sheet!Row,Col) are recorded."""
        excel_bytes = self._create_sample_excel(
            headers=["Shipper Location", "Consignee Location", "Gross Weight lbs"],
            rows_data=[
                ["CALGARY, AB", "LINDSAY, ON", 1450]
            ]
        )

        resp = self.client.post(
            "/api/quotes/analyze-excel",
            files={"file": ("trace_test.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        row = data["preview_rows"][0]

        # Coordinates must contain real column letters and row indices
        self.assertIn("Shipments!A2", row["origin_coord"])
        self.assertIn("Shipments!B2", row["dest_coord"])
        self.assertIn("Shipments!C2", row["weight_coord"])

    def test_03_rule5_metric_unit_and_currency_detection(self):
        """Rule 5: Detect explicit metric units (kg, cm) and currency without guessing."""
        excel_bytes = self._create_sample_excel(
            headers=["Pickup City", "Delivery City", "Total Weight (kg)", "Dims (cm)", "Declared Value (CAD)"],
            rows_data=[
                ["CALGARY, AB", "LINDSAY, ON", 500, "100x100x100", 2500]
            ]
        )

        resp = self.client.post(
            "/api/quotes/analyze-excel",
            files={"file": ("metric_test.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        units = data["detected_units"]

        self.assertEqual(units["weight_unit"], "kg")
        self.assertEqual(units["dimension_unit"], "cm")
        self.assertEqual(units["currency"], "CAD")

        # 500 kg converted to ~1,102.31 lbs
        preview = data["preview_rows"][0]
        self.assertAlmostEqual(preview["weight_lbs"], 1102.31, places=1)

    def test_04_download_reformatted_excel_file(self):
        """Verifies GET /api/quotes/download-reformatted/{token} produces a valid formatted workbook."""
        excel_bytes = self._create_sample_excel(
            headers=["Origin", "Destination", "Weight"],
            rows_data=[
                ["CALGARY, AB", "LINDSAY, ON", 1450]
            ]
        )

        # Analyze
        resp = self.client.post(
            "/api/quotes/analyze-excel",
            files={"file": ("download_test.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        token = resp.json()["analysis_token"]

        # Confirm
        self.client.post("/api/quotes/confirm-reformatted", json={
            "analysis_token": token,
            "row_corrections": {}
        })

        # Download
        down_resp = self.client.get(f"/api/quotes/download-reformatted/{token}")
        self.assertEqual(down_resp.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", down_resp.headers["content-type"])

        # Validate downloaded workbook structure
        down_wb = openpyxl.load_workbook(io.BytesIO(down_resp.content))
        self.assertIn("Normalized Shipments", down_wb.sheetnames)
        self.assertIn("Extraction Audit Log", down_wb.sheetnames)

        ws_norm = down_wb["Normalized Shipments"]
        self.assertEqual(ws_norm.cell(row=1, column=1).value, "Row ID")
        self.assertEqual(ws_norm.cell(row=2, column=2).value, "CALGARY, AB")

    def test_05_flag_exceptions_as_non_critical_info_and_ignore_fields(self):
        """
        User Requirement: Flag raised exceptions as non-critical info (e.g. personal business info
        such as company headquarters or internal dispatch) to allow agent to ignore fields
        without forcing strict value-based input before quoting.
        """
        # Sheet has missing origin (implied company headquarters) and missing weight
        excel_bytes = self._create_sample_excel(
            headers=["Origin", "Destination", "Weight"],
            rows_data=[
                [None, "LINDSAY, ON", None]  # Row 2 in Excel: missing origin & weight
            ]
        )

        resp = self.client.post(
            "/api/quotes/analyze-excel",
            files={"file": ("company_hq_order.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        token = data["analysis_token"]
        self.assertTrue(data["has_blockers"])

        # Without flagging or correcting, confirmation must be blocked by Rule 25
        fail_resp = self.client.post("/api/quotes/confirm-reformatted", json={
            "analysis_token": token,
            "row_corrections": {}
        })
        self.assertEqual(fail_resp.status_code, 400)
        self.assertIn("Rule 25 Requirement", fail_resp.json()["detail"])

        # With user flagging the exception as non-critical info (e.g. Company HQ / internal dispatch):
        # Strict value-based input is NOT forced, agent ignores the missing fields, and quoting succeeds!
        success_resp = self.client.post("/api/quotes/confirm-reformatted", json={
            "analysis_token": token,
            "non_critical_rows": [2],
            "row_corrections": {}  # NO manual text typed!
        })
        self.assertEqual(success_resp.status_code, 200)
        res_data = success_resp.json()
        self.assertEqual(res_data["status"], "success")
        self.assertEqual(res_data["total_shipments"], 1)

        # Download the reformatted workbook and verify non-critical status and audit entries
        down_resp = self.client.get(f"/api/quotes/download-reformatted/{token}")
        self.assertEqual(down_resp.status_code, 200)
        down_wb = openpyxl.load_workbook(io.BytesIO(down_resp.content))

        ws_norm = down_wb["Normalized Shipments"]
        # Row 2 Status must be NON_CRITICAL_INFO
        self.assertEqual(ws_norm.cell(row=2, column=9).value, "NON_CRITICAL_INFO")

        ws_audit = down_wb["Extraction Audit Log"]
        audit_rows = list(ws_audit.iter_rows(values_only=True))
        # Ensure audit log captured the non-critical bypass
        bypass_logged = any(r[3] == "NON_CRITICAL_BYPASS" for r in audit_rows)
        self.assertTrue(bypass_logged, "Extraction Audit Log must record non-critical bypass")

    def test_06_unmapped_company_headquarters_column_flagged_as_non_critical(self):
        """Verifies unmapped business columns (e.g. Company Headquarters, PO, Notes) are detected and ignored as non-critical."""
        excel_bytes = self._create_sample_excel(
            headers=["Company Headquarters", "Origin", "Destination", "Weight", "Sales Rep Code"],
            rows_data=[
                ["Headquarters Bldg 4, 100 Main St", "CALGARY, AB", "LINDSAY, ON", 1200, "SR-994"]
            ]
        )

        resp = self.client.post(
            "/api/quotes/analyze-excel",
            files={"file": ("unmapped_columns_test.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()

        # Unmapped columns must identify 'Company Headquarters' and 'Sales Rep Code'
        unmapped_headers = [c["header"] for c in data.get("unmapped_columns", [])]
        self.assertIn("Company Headquarters", unmapped_headers)
        self.assertIn("Sales Rep Code", unmapped_headers)

        # Confirming with non_critical_columns succeeds cleanly
        token = data["analysis_token"]
        conf_resp = self.client.post("/api/quotes/confirm-reformatted", json={
            "analysis_token": token,
            "non_critical_columns": ["Company Headquarters", "Sales Rep Code"],
            "flag_all_exceptions_non_critical": True
        })
        self.assertEqual(conf_resp.status_code, 200)
        self.assertEqual(conf_resp.json()["status"], "success")

if __name__ == "__main__":
    unittest.main()
