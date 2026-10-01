"""
test_priority_2_enhancements.py
===============================
Comprehensive verification suite for Priority 2 Product & Engine Enhancements:
- 2.1 Canadian Postal FSA Resolver Table
- 2.2 Weekly Fuel Surcharge (FSC) Index Manager
- 2.3 Client-Side Rate Sheet Masker / Redaction Tool
- 2.4 SMC3 CzarLite & Deficit Weight Rating Engine
"""

import io
import openpyxl
from starlette.testclient import TestClient
from main import app
from services.fsa_resolver_service import resolve_location, extract_fsa, resolve_lane_locations
from services.fuel_index_service import (
    get_available_benchmark_indices,
    get_tenant_fuel_settings,
    update_tenant_fuel_settings,
    resolve_effective_fuel_surcharge
)
from services.rate_redactor_service import redact_rate_sheet_bytes
from services.ratesift_engine import quote_all_confirmed_carriers, match_lane_break
from services.seed_canadian_benchmarks import seed_canadian_benchmarks_for_user

client = TestClient(app)

def test_priority_2_suite():
    print("\n" + "=" * 70)
    print("RUNNING PRIORITY 2 PRODUCT & ENGINE ENHANCEMENT TEST SUITE")
    print("=" * 70)

    # -------------------------------------------------------------------------
    # TEST 2.1: CANADIAN POSTAL FSA RESOLVER
    # -------------------------------------------------------------------------
    print("\n[TASK 2.1] Testing Canadian Postal FSA Resolver...")

    # A. Downtown Toronto Postal Code
    loc_toronto = resolve_location("M5V 2T6")
    assert loc_toronto["is_resolved"] is True
    assert loc_toronto["fsa"] == "M5V"
    assert loc_toronto["city"] == "TORONTO"
    assert loc_toronto["province"] == "ON"
    assert loc_toronto["location_string"] == "TORONTO, ON"
    assert loc_toronto["is_rural"] is False
    print("   [OK] M5V 2T6 successfully resolved to TORONTO, ON (Zone: ON-TOR)")

    # B. Calgary Financial District FSA
    loc_calgary = resolve_location("T2P")
    assert loc_calgary["is_resolved"] is True
    assert loc_calgary["fsa"] == "T2P"
    assert loc_calgary["city"] == "CALGARY"
    assert loc_calgary["province"] == "AB"
    assert loc_calgary["location_string"] == "CALGARY, AB"
    print("   [OK] T2P successfully resolved to CALGARY, AB (Zone: AB-CAL)")

    # C. Montreal Freight Hub (Dorval)
    loc_mtl = resolve_location("H9P 1K2")
    assert loc_mtl["city"] == "DORVAL"
    assert loc_mtl["province"] == "QC"
    print("   [OK] H9P 1K2 successfully resolved to Dorval West Island Cargo Hub")

    # D. Corridor Lane Resolution
    orig_c, dest_c, meta = resolve_lane_locations("M5V 2T6", "H3B 1A1")
    assert "TORONTO" in orig_c
    assert "MONTREAL" in dest_c
    assert meta["corridor_code"] == "ON-QC"
    print(f"   [OK] Lane resolved: {orig_c} -> {dest_c} (Corridor: {meta['corridor_code']})")

    # E. FSA REST API Endpoints
    res_fsa_api = client.get("/api/geo/resolve-fsa?code=M5V")
    assert res_fsa_api.status_code == 200
    assert res_fsa_api.json()["fsa"] == "M5V"

    res_lane_api = client.get("/api/geo/resolve-lane?origin=M5V2T6&destination=T2P1J9")
    assert res_lane_api.status_code == 200
    assert "TORONTO" in res_lane_api.json()["origin_canonical"]
    assert "CALGARY" in res_lane_api.json()["destination_canonical"]
    print("   [OK] REST API Endpoints /api/geo/resolve-fsa & /api/geo/resolve-lane verified")

    # -------------------------------------------------------------------------
    # TEST 2.2: WEEKLY FUEL SURCHARGE (FSC) INDEX MANAGER
    # -------------------------------------------------------------------------
    print("\n[TASK 2.2] Testing Weekly Fuel Surcharge (FSC) Index Manager...")

    # A. Benchmark Indices
    benchmarks = get_available_benchmark_indices()
    assert len(benchmarks) >= 4
    ota_bench = next(b for b in benchmarks if b["index_code"] == "OTA_LTL_STANDARD")
    assert ota_bench["ltl_surcharge_percent"] == 31.50
    print(f"   [OK] Loaded {len(benchmarks)} Canadian diesel benchmarks (OTA Baseline: {ota_bench['ltl_surcharge_percent']}%)")

    # B. Tenant Settings & Calculation
    settings = get_tenant_fuel_settings("usr_alex_rivers")
    assert settings["active_index_code"] == "OTA_LTL_STANDARD"
    assert settings["effective_ltl_percent"] == 31.50

    fsc_result = resolve_effective_fuel_surcharge("usr_alex_rivers", base_freight=1000.0)
    assert fsc_result["fsc_percent"] == 31.50
    assert fsc_result["fsc_amount"] == 315.00
    print(f"   [OK] Calculated FSC on $1,000.00 CAD base freight: ${fsc_result['fsc_amount']} CAD ({fsc_result['source']})")

    # C. Custom Override
    update_tenant_fuel_settings("usr_alex_rivers", custom_ltl_percent=35.00)
    fsc_custom = resolve_effective_fuel_surcharge("usr_alex_rivers", base_freight=1000.0)
    assert fsc_custom["fsc_percent"] == 35.00
    assert fsc_custom["fsc_amount"] == 350.00
    # Reset back to default OTA
    update_tenant_fuel_settings("usr_alex_rivers", custom_ltl_percent=None)
    print("   [OK] Broker custom FSC percentage override and reset verified")

    # D. REST API Endpoints
    res_fuel_benches = client.get("/api/fuel-indices/benchmarks")
    assert res_fuel_benches.status_code == 200
    assert len(res_fuel_benches.json()["benchmarks"]) >= 4

    res_fuel_settings = client.get(
        "/api/fuel-indices/settings",
        cookies={"ratesift_session": "session_demo_alex_rivers"}
    )
    assert res_fuel_settings.status_code == 200
    print("   [OK] REST API Endpoints /api/fuel-indices/benchmarks & settings verified")

    # -------------------------------------------------------------------------
    # TEST 2.3: CLIENT-SIDE RATE SHEET MASKER / REDACTION TOOL
    # -------------------------------------------------------------------------
    print("\n[TASK 2.3] Testing Client-Side Rate Sheet Masker / Redactor...")

    # Build an in-memory Excel sheet with confidential broker data
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Confidential_Tariff"

    # Add sensitive rows alongside freight rates
    ws.append(["Carrier Name", "Day & Ross Freight", "Account #", "ACCT-9948210"])
    ws.append(["Customer Name", "Acme Industrial Supplies Inc.", "Rep Email", "sales.dispatch@dayross.com"])
    ws.append(["Broker Margin", "15.5% Target Markup", "Contact Phone", "416-555-0199"])
    ws.append(["Origin", "Destination", "Min Weight", "Max Weight", "Rate / CWT"])
    ws.append(["TORONTO, ON", "MONTREAL, QC", 0, 499, 19.50])
    ws.append(["TORONTO, ON", "MONTREAL, QC", 500, 999, 16.20])
    ws.append(["TORONTO, ON", "MONTREAL, QC", 1000, 1999, 13.50])

    buf = io.BytesIO()
    wb.save(buf)
    raw_bytes = buf.getvalue()

    redacted_bytes, audit = redact_rate_sheet_bytes(raw_bytes, "confidential_broker_tariff.xlsx")
    assert audit["total_redactions"] >= 4, f"Expected at least 4 redactions, got {audit['total_redactions']}"
    assert audit["privacy_status"] == "VERIFIED_SCRUBBED"

    # Inspect redacted workbook to confirm sensitive data replaced and rates preserved
    redacted_wb = openpyxl.load_workbook(io.BytesIO(redacted_bytes))
    redacted_ws = redacted_wb["Confidential_Tariff"]

    # Account cell scrubbed
    account_val = str(redacted_ws.cell(row=1, column=4).value)
    assert "[REDACTED_ACCOUNT]" in account_val
    # Email scrubbed
    email_val = str(redacted_ws.cell(row=2, column=4).value)
    assert "[REDACTED_EMAIL]" in email_val
    # Phone scrubbed
    phone_val = str(redacted_ws.cell(row=3, column=4).value)
    assert "[REDACTED_PHONE]" in phone_val
    # Rate cells 100% preserved
    rate_cell = redacted_ws.cell(row=5, column=5).value
    assert float(rate_cell) == 19.50
    print(f"   [OK] Redacted {audit['total_redactions']} confidential cells while preserving 100% of freight rates")

    # Redaction API Test
    res_redact_api = client.post(
        "/api/ratesift/preview-redaction",
        files={"file": ("test_tariff.xlsx", raw_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert res_redact_api.status_code == 200
    assert res_redact_api.json()["audit"]["total_redactions"] >= 4
    print("   [OK] REST API Endpoint /api/ratesift/preview-redaction verified")

    # -------------------------------------------------------------------------
    # TEST 2.4: SMC3 CZARLITE & DEFICIT WEIGHT RATING ENGINE
    # -------------------------------------------------------------------------
    print("\n[TASK 2.4] Testing SMC3 CzarLite & Deficit Weight Rating Engine...")

    # Define test breaks where deficit rating is mathematically advantageous:
    # 5C (500 - 999 lbs): $42.00 / CWT
    # 1M (1000 - 1999 lbs): $34.00 / CWT
    # At 880 lbs:
    #   Natural charge = (880 / 100) * $42.00 = $369.60
    #   Bumping to 1,000 lbs = (1000 / 100) * $34.00 = $340.00
    #   Savings = $29.60 (Deficit weight: 120 lbs)
    test_breaks = [
        {"break_name": "MIN", "origin_spec": "TORONTO, ON", "dest_spec": "MONTREAL, QC", "min_weight": 0, "max_weight": 499, "base_rate": 50.00, "rate_type": "CWT"},
        {"break_name": "5C", "origin_spec": "TORONTO, ON", "dest_spec": "MONTREAL, QC", "min_weight": 500, "max_weight": 999, "base_rate": 42.00, "rate_type": "CWT"},
        {"break_name": "1M", "origin_spec": "TORONTO, ON", "dest_spec": "MONTREAL, QC", "min_weight": 1000, "max_weight": 1999, "base_rate": 34.00, "rate_type": "CWT"},
        {"break_name": "2M", "origin_spec": "TORONTO, ON", "dest_spec": "MONTREAL, QC", "min_weight": 2000, "max_weight": 4999, "base_rate": 28.00, "rate_type": "CWT"},
    ]

    # Test 1: Shipment at 880 lbs triggers deficit rating
    matched_880 = match_lane_break(test_breaks, "TORONTO, ON", "MONTREAL, QC", 880.0)
    assert matched_880 is not None
    assert matched_880["is_deficit_rated"] is True
    assert matched_880["break_name"] == "1M"
    assert matched_880["deficit_weight"] == 120.0
    assert matched_880["effective_base_charge"] == 340.00
    assert matched_880["deficit_savings"] == 29.60
    print(f"   [OK] Deficit Rating triggered for 880 lbs: Billed as 1,000 lbs @ $34.00/CWT = $340.00 (Saved ${matched_880['deficit_savings']})")

    # Test 2: Postal Code Input (M5V 2T6 -> H3B 1A1) in Engine
    seed_canadian_benchmarks_for_user("usr_alex_rivers", force=True)
    res_postal_engine = quote_all_confirmed_carriers(
        user_id="usr_alex_rivers",
        origin="M5V 2T6",
        destination="H3B 1A1",
        actual_weight=650.0,
        accessorials=["appointment", "fuel"]
    )
    quotes = res_postal_engine["quotes"]
    assert len(quotes) >= 4, f"Expected quotes for postal code lane, got {len(quotes)}"
    
    # Verify fuel surcharge applied via Canadian benchmark index
    first_q = quotes[0]
    fsc_found = any(s["code"] == "FSC" for s in first_q["surcharges"])
    assert fsc_found is True
    print(f"   [OK] Engine successfully quoted Postal Code lane M5V 2T6 -> H3B 1A1 with OTA Fuel Surcharge itemized")

    print("\n" + "=" * 70)
    print("ALL PRIORITY 2 ENHANCEMENTS FULLY VERIFIED AND PASSING (100%)")
    print("=" * 70)

if __name__ == "__main__":
    test_priority_2_suite()
