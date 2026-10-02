import io
import uuid
from datetime import datetime, timedelta
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_client_proposal_data(
    quote_data: Dict[str, Any],
    client_name: str = "Client Partner",
    broker_info: Optional[Dict[str, Any]] = None,
    markup_pct: float = 15.0,
    custom_total: Optional[float] = None
) -> Dict[str, Any]:
    """
    Constructs normalized client proposal data with broker margin markups.
    Hides internal carrier wholesale net rates and presents professional all-in pricing.
    """
    broker = broker_info or {}
    company_name = broker.get("company") or "TechCorp Logistics"
    agent_name = broker.get("name") or "Alex Rivers"
    agent_email = broker.get("email") or "alex.rivers@techcorp.io"
    origin_hub = broker.get("origin_zip") or quote_data.get("origin") or "TORONTO, ON"

    now = datetime.utcnow()
    proposal_id = f"PROP-{now.strftime('%Y%m%d')}-{uuid.uuid4().hex[:5].upper()}"
    expiry_date = (now + timedelta(days=7)).strftime("%B %d, %Y")
    issue_date = now.strftime("%B %d, %Y")

    base_cost = float(quote_data.get("base_rate", 0.0))
    surcharges = quote_data.get("surcharges", [])
    total_surcharges = float(quote_data.get("total_surcharges", 0.0))
    net_carrier_total = float(quote_data.get("final_total", base_cost + total_surcharges))

    # Calculate client pricing
    if custom_total is not None and custom_total > 0:
        client_total = round(float(custom_total), 2)
        client_base_freight = round(max(0.0, client_total - total_surcharges), 2)
        effective_markup = round(((client_total - net_carrier_total) / net_carrier_total) * 100.0, 1) if net_carrier_total > 0 else 0.0
    else:
        markup_multiplier = 1.0 + (max(0.0, markup_pct) / 100.0)
        client_base_freight = round(base_cost * markup_multiplier, 2)
        client_total = round(client_base_freight + total_surcharges, 2)
        effective_markup = round(markup_pct, 1)

    margin_dollars = round(client_total - net_carrier_total, 2)

    # Line items for client view
    line_items = [
        {
            "category": "Freight Transportation",
            "description": f"Standard LTL Road Service ({quote_data.get('origin', '')} → {quote_data.get('destination', '')})",
            "amount": client_base_freight
        }
    ]

    for s in surcharges:
        line_items.append({
            "category": "Accessorial Service",
            "description": s.get("name") or s.get("surcharge_code") or "Additional Service",
            "amount": float(s.get("amount", 0.0))
        })

    proposal_dict = {
        "proposal_id": proposal_id,
        "issue_date": issue_date,
        "expiry_date": expiry_date,
        "broker": {
            "company": company_name,
            "agent_name": agent_name,
            "email": agent_email,
            "origin_hub": origin_hub
        },
        "client": {
            "name": client_name.strip() if client_name else "Valued Client",
            "attention": "Freight Logistics & Shipping Dept"
        },
        "shipment": {
            "origin": quote_data.get("origin", "N/A"),
            "destination": quote_data.get("destination", "N/A"),
            "weight_lbs": quote_data.get("weight_lbs") or quote_data.get("actual_weight") or quote_data.get("billable_weight") or 0.0,
            "billable_weight": quote_data.get("billable_weight") or quote_data.get("weight_lbs") or quote_data.get("actual_weight") or 0.0,
            "is_dim_billed": bool(quote_data.get("is_dim_billed", False)),
            "service_level": quote_data.get("service_name") or "Standard Road LTL",
            "transit_days": quote_data.get("transit_days") or 3,
            "currency": quote_data.get("currency") or "CAD"
        },
        "pricing": {
            "client_base_freight": client_base_freight,
            "total_surcharges": total_surcharges,
            "client_total": client_total,
            "currency": quote_data.get("currency") or "CAD",
            "effective_markup_pct": effective_markup,
            "broker_margin_dollars": margin_dollars
        },
        "line_items": line_items,
        "terms": [
            "1. Rate Validity: This quote is valid for 7 calendar days from issuance and is subject to equipment availability.",
            "2. Inspection & Verification: Final billing is subject to carrier re-weigh, cube, and standard freight class verification.",
            "3. Transit Times: Transit days are business-day estimates and exclude statutory holidays and weekend layovers.",
            "4. Loading / Unloading: Standard 2-hour free time applies at shipper and receiver docks. Detention applies thereafter."
        ]
    }
    proposal_dict["email"] = generate_proposal_email(proposal_dict)
    return proposal_dict

def generate_proposal_email(proposal_data: Dict[str, Any]) -> Dict[str, str]:
    """
    Generates a professional, executive-grade freight proposal email with full quote details.
    Includes route specifications, itemized pricing breakdown, rate lock, and commercial terms.
    """
    broker = proposal_data.get("broker", {})
    client = proposal_data.get("client", {})
    shipment = proposal_data.get("shipment", {})
    pricing = proposal_data.get("pricing", {})
    line_items = proposal_data.get("line_items", [])

    origin = shipment.get("origin", "N/A")
    destination = shipment.get("destination", "N/A")
    weight = float(shipment.get("billable_weight") or shipment.get("weight_lbs") or 0.0)
    service_level = shipment.get("service_level", "Standard Road LTL")
    transit_days = shipment.get("transit_days", 3)
    currency = pricing.get("currency", "CAD")
    client_total = float(pricing.get("client_total", 0.0))
    proposal_id = proposal_data.get("proposal_id", "PROP-REF")
    issue_date = proposal_data.get("issue_date", "")
    expiry_date = proposal_data.get("expiry_date", "")
    client_name = client.get("name", "Client Partner")
    agent_name = broker.get("agent_name", "Alex Rivers")
    company_name = broker.get("company", "TechCorp Logistics")
    agent_email = broker.get("email", "alex.rivers@techcorp.io")
    origin_hub = broker.get("origin_hub", "Toronto, ON")

    subject = f"Freight Rate Quote Proposal: {origin} → {destination} ({weight:,.0f} lbs) - Ref #{proposal_id}"

    # Build itemized lines for text & HTML
    item_lines_txt = []
    item_rows_html = []
    for item in line_items:
        desc = item.get("description", "Transportation Charge")
        cat = item.get("category", "Service")
        amt = float(item.get("amount", 0.0))
        item_lines_txt.append(f"  • {desc} ({cat}): ${amt:,.2f} {currency}")
        item_rows_html.append(f"""
          <tr>
            <td style="padding: 8px 12px; border-bottom: 1px solid #e2e8f0; font-size: 13px; color: #1e293b;">{desc}</td>
            <td style="padding: 8px 12px; border-bottom: 1px solid #e2e8f0; font-size: 12px; color: #64748b;">{cat}</td>
            <td style="padding: 8px 12px; border-bottom: 1px solid #e2e8f0; font-size: 13px; font-weight: bold; color: #0f172a; text-align: right; font-family: monospace;">${amt:,.2f}</td>
          </tr>
        """)

    items_block = "\n".join(item_lines_txt)
    items_table_html = "".join(item_rows_html)

    body_text = f"""Hi {client_name},

Thank you for the opportunity to quote your freight shipment. Please find your official freight rate proposal and shipment details below:

============================================================
SHIPMENT OVERVIEW & ROUTE PROFILE
============================================================
• Origin:             {origin}
• Destination:        {destination}
• Service Level:      {service_level}
• Estimated Transit:  {transit_days} Business Days
• Cargo Weight:       {weight:,.1f} lbs
• Proposal Reference: {proposal_id}
• Date Issued:        {issue_date}
• Rate Lock Guarantee: 7-Day Rate Lock Guarantee (Valid until {expiry_date})

============================================================
ITEMIZED RATE SCHEDULE ({currency})
============================================================
{items_block}
------------------------------------------------------------
GUARANTEED ALL-IN CLIENT TOTAL: ${client_total:,.2f} {currency}
------------------------------------------------------------

TERMS & CONDITIONS OF CARRIAGE:
1. Rate Guarantee: Guaranteed valid for 7 calendar days from issuance. Subject to carrier equipment availability.
2. Verification: Final billing subject to standard scale re-weigh and cube audit upon tender.
3. Transit Times: Business day estimates excluding statutory holidays and weekends.
4. Loading / Unloading: Standard 2-hour free time applies at shipper and receiver docks.

To accept this quote and schedule dispatch pickup, simply reply directly to this email or call our operations desk.

Best regards,

{agent_name}
{company_name}
Email: {agent_email}
Dispatch Hub: {origin_hub}
RateSift Canadian Freight Engine • WHC Canadian Data Residency Verified"""

    body_html = f"""<div style="font-family: Arial, -apple-system, sans-serif; color: #1e293b; max-width: 650px; line-height: 1.6;">
  <p style="font-size: 14px; margin-bottom: 16px;">Hi <strong>{client_name}</strong>,</p>
  <p style="font-size: 13px; color: #475569; margin-bottom: 20px;">
    Thank you for the opportunity to quote your freight shipment. Please find your formal rate proposal and route details below:
  </p>

  <div style="background-color: #f8fafc; border: 1px solid #e2e8f0; border-radius: 8px; padding: 16px; margin-bottom: 20px;">
    <div style="font-size: 11px; font-weight: bold; color: #64748b; text-transform: uppercase; margin-bottom: 12px; letter-spacing: 0.5px;">
      Shipment Specifications & Route Profile
    </div>
    <table style="width: 100%; border-collapse: collapse; font-size: 13px;">
      <tr>
        <td style="padding: 4px 0; color: #64748b; width: 35%;">Origin:</td>
        <td style="padding: 4px 0; font-weight: bold; color: #0f172a;">{origin}</td>
      </tr>
      <tr>
        <td style="padding: 4px 0; color: #64748b;">Destination:</td>
        <td style="padding: 4px 0; font-weight: bold; color: #0f172a;">{destination}</td>
      </tr>
      <tr>
        <td style="padding: 4px 0; color: #64748b;">Service Level:</td>
        <td style="padding: 4px 0; font-weight: bold; color: #0f172a;">{service_level}</td>
      </tr>
      <tr>
        <td style="padding: 4px 0; color: #64748b;">Estimated Transit:</td>
        <td style="padding: 4px 0; font-weight: bold; color: #0f172a;">{transit_days} Business Days</td>
      </tr>
      <tr>
        <td style="padding: 4px 0; color: #64748b;">Billable Weight:</td>
        <td style="padding: 4px 0; font-weight: bold; color: #0f172a;">{weight:,.1f} lbs</td>
      </tr>
      <tr>
        <td style="padding: 4px 0; color: #64748b;">Proposal Ref:</td>
        <td style="padding: 4px 0; font-family: monospace; color: #dc2626; font-weight: bold;">{proposal_id}</td>
      </tr>
      <tr>
        <td style="padding: 4px 0; color: #64748b;">Rate Lock Guarantee:</td>
        <td style="padding: 4px 0; color: #16a34a; font-weight: bold;">Valid until {expiry_date}</td>
      </tr>
    </table>
  </div>

  <div style="border: 1px solid #e2e8f0; border-radius: 8px; overflow: hidden; margin-bottom: 20px;">
    <table style="width: 100%; border-collapse: collapse;">
      <thead>
        <tr style="background-color: #0f172a; color: white;">
          <th style="padding: 10px 12px; font-size: 11px; text-align: left; text-transform: uppercase;">Description</th>
          <th style="padding: 10px 12px; font-size: 11px; text-align: left; text-transform: uppercase;">Category</th>
          <th style="padding: 10px 12px; font-size: 11px; text-align: right; text-transform: uppercase;">Amount ({currency})</th>
        </tr>
      </thead>
      <tbody>
        {items_table_html}
      </tbody>
      <tfoot>
        <tr style="background-color: #fef2f2; border-top: 2px solid #dc2626;">
          <td colspan="2" style="padding: 12px; font-size: 13px; font-weight: bold; color: #991b1b; text-transform: uppercase;">Total Client Price (All-In {currency})</td>
          <td style="padding: 12px; font-size: 16px; font-weight: 900; color: #dc2626; text-align: right; font-family: monospace;">${client_total:,.2f} {currency}</td>
        </tr>
      </tfoot>
    </table>
  </div>

  <div style="font-size: 11px; color: #64748b; line-height: 1.5; margin-bottom: 24px; padding-left: 8px; border-left: 3px solid #cbd5e1;">
    <p style="margin: 0 0 4px 0;"><strong>Terms:</strong> Quote valid for 7 calendar days. Final billing subject to actual scale re-weigh and dimensional cube inspection.</p>
    <p style="margin: 0;">Standard 2-hour dock free time applies for loading and unloading.</p>
  </div>

  <p style="font-size: 13px; margin-bottom: 24px;">
    To accept this quote and schedule dispatch pickup, please reply directly to this email.
  </p>

  <div style="border-top: 1px solid #e2e8f0; padding-top: 16px; font-size: 12px; color: #334155;">
    <strong style="font-size: 13px; color: #0f172a;">{agent_name}</strong><br>
    <span>{company_name}</span><br>
    <span>Email: <a href="mailto:{agent_email}" style="color: #dc2626; text-decoration: none;">{agent_email}</a></span><br>
    <span>Dispatch Hub: {origin_hub}</span>
  </div>
</div>"""

    return {
        "subject": subject,
        "body_text": body_text,
        "body_html": body_html
    }

def generate_proposal_excel_workbook(proposal_data: Dict[str, Any]) -> bytes:
    """
    Renders an executive-grade Excel quote proposal (.xlsx) using openpyxl.
    Features professional typography, navy header, formatted currency, and clean printable layout.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Client_Freight_Proposal"

    # Configure grid lines and page layout
    ws.views.sheetView[0].showGridLines = True
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    # Palette definitions
    NAVY_HEADER = "0F172A"
    ICE_BLUE = "F1F5F9"
    HIGHLIGHT_RED = "DC2626"
    BORDER_GRAY = "CBD5E1"
    ZEBRA_LIGHT = "F8FAFC"

    font_title = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Arial", size=9, bold=False, color="94A3B8")
    font_section_hdr = Font(name="Arial", size=10, bold=True, color="0F172A")
    font_label = Font(name="Arial", size=9, bold=True, color="64748B")
    font_val = Font(name="Arial", size=10, bold=False, color="0F172A")
    font_val_bold = Font(name="Arial", size=10, bold=True, color="0F172A")
    font_table_hdr = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    font_total_label = Font(name="Arial", size=11, bold=True, color="0F172A")
    font_total_val = Font(name="Arial", size=14, bold=True, color="DC2626")
    font_terms = Font(name="Arial", size=8, italic=True, color="64748B")

    fill_header = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    fill_subhdr = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_ice = PatternFill(start_color=ICE_BLUE, end_color=ICE_BLUE, fill_type="solid")
    fill_zebra = PatternFill(start_color=ZEBRA_LIGHT, end_color=ZEBRA_LIGHT, fill_type="solid")
    fill_total = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color=BORDER_GRAY),
        right=Side(style="thin", color=BORDER_GRAY),
        top=Side(style="thin", color=BORDER_GRAY),
        bottom=Side(style="thin", color=BORDER_GRAY)
    )
    thick_bottom = Border(bottom=Side(style="medium", color=NAVY_HEADER))
    double_bottom = Border(
        top=Side(style="thin", color=BORDER_GRAY),
        bottom=Side(style="double", color="DC2626")
    )

    # 1. Header Banner (Rows 1-3)
    ws.merge_cells("A1:E1")
    ws["A1"] = f"{proposal_data['broker']['company'].upper()}"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_header
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells("A2:E2")
    ws["A2"] = f"FREIGHT RATE QUOTE PROPOSAL  •  REF: {proposal_data['proposal_id']}"
    ws["A2"].font = Font(name="Arial", size=10, bold=True, color="F87171")
    ws["A2"].fill = fill_header
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells("A3:E3")
    ws["A3"] = f"Issued: {proposal_data['issue_date']}   |   Valid Until: {proposal_data['expiry_date']} (7-Day Rate Lock)"
    ws["A3"].font = font_subtitle
    ws["A3"].fill = fill_header
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 10

    # 2. Broker & Client Information Cards (Rows 5-8)
    ws["A5"] = "PREPARED FOR (CLIENT)"
    ws["A5"].font = font_section_hdr
    ws["D5"] = "LOGISTICS DISPATCH / BROKER"
    ws["D5"].font = font_section_hdr

    client_info = [
        ("Client Account:", proposal_data["client"]["name"]),
        ("Attention:", proposal_data["client"]["attention"]),
        ("Status:", "Standard Commercial Terms")
    ]
    for idx, (label, val) in enumerate(client_info, start=6):
        ws[f"A{idx}"] = label
        ws[f"A{idx}"].font = font_label
        ws[f"B{idx}"] = val
        ws[f"B{idx}"].font = font_val_bold if idx == 6 else font_val

    broker_info = [
        ("Prepared By:", proposal_data["broker"]["agent_name"]),
        ("Contact Email:", proposal_data["broker"]["email"]),
        ("Dispatch Hub:", proposal_data["broker"]["origin_hub"])
    ]
    for idx, (label, val) in enumerate(broker_info, start=6):
        ws[f"D{idx}"] = label
        ws[f"D{idx}"].font = font_label
        ws[f"E{idx}"] = val
        ws[f"E{idx}"].font = font_val_bold if idx == 6 else font_val

    ws.row_dimensions[9].height = 10

    # 3. Shipment Specifications Card (Rows 10-14)
    ws.merge_cells("A10:E10")
    ws["A10"] = "SHIPMENT SPECIFICATIONS & ROUTE PROFILE"
    ws["A10"].font = Font(name="Arial", size=9, bold=True, color="0F172A")
    ws["A10"].fill = fill_ice
    ws["A10"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    specs = [
        ("Origin Lane:", proposal_data["shipment"]["origin"], "Mode / Service:", proposal_data["shipment"]["service_level"]),
        ("Destination Lane:", proposal_data["shipment"]["destination"], "Transit Estimate:", f"{proposal_data['shipment']['transit_days']} Business Days"),
        ("Billable Weight:", f"{proposal_data['shipment']['billable_weight']:,.1f} lbs", "Currency:", f"{proposal_data['shipment']['currency']} ($)")
    ]

    for idx, (l1, v1, l2, v2) in enumerate(specs, start=11):
        ws[f"A{idx}"] = l1
        ws[f"A{idx}"].font = font_label
        ws[f"B{idx}"] = v1
        ws[f"B{idx}"].font = font_val_bold
        ws[f"D{idx}"] = l2
        ws[f"D{idx}"].font = font_label
        ws[f"E{idx}"] = v2
        ws[f"E{idx}"].font = font_val_bold
        ws.row_dimensions[idx].height = 18

    ws.row_dimensions[14].height = 10

    # 4. Itemized Quotation Table (Rows 15+)
    start_table_row = 15
    table_headers = ["ITEM #", "CHARGE CATEGORY", "DESCRIPTION / DETAILS", "CURRENCY", "AMOUNT"]
    cols = ["A", "B", "C", "D", "E"]

    for col, h in zip(cols, table_headers):
        cell = ws[f"{col}{start_table_row}"]
        cell.value = h
        cell.font = font_table_hdr
        cell.fill = fill_subhdr
        cell.alignment = Alignment(horizontal="center" if col in ["A", "D"] else ("right" if col == "E" else "left"), vertical="center")

    ws.row_dimensions[start_table_row].height = 20

    current_r = start_table_row + 1
    for i, item in enumerate(proposal_data["line_items"], start=1):
        ws[f"A{current_r}"] = f"{i:02d}"
        ws[f"B{current_r}"] = item["category"]
        ws[f"C{current_r}"] = item["description"]
        ws[f"D{current_r}"] = proposal_data["pricing"]["currency"]
        ws[f"E{current_r}"] = item["amount"]

        ws[f"A{current_r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"D{current_r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"E{current_r}"].alignment = Alignment(horizontal="right", vertical="center")
        ws[f"E{current_r}"].number_format = "$#,##0.00"

        # Apply zebra shading
        row_fill = fill_zebra if i % 2 == 0 else PatternFill(fill_type=None)
        for col in cols:
            c = ws[f"{col}{current_r}"]
            c.font = font_val
            c.border = thin_border
            if row_fill.fill_type:
                c.fill = row_fill

        ws.row_dimensions[current_r].height = 20
        current_r += 1

    # Total Box
    ws.merge_cells(f"A{current_r}:C{current_r}")
    ws[f"A{current_r}"] = "TOTAL CLIENT QUOTED RATE (ALL-IN CAD)"
    ws[f"A{current_r}"].font = font_total_label
    ws[f"A{current_r}"].alignment = Alignment(horizontal="right", vertical="center")
    ws[f"A{current_r}"].fill = fill_total

    ws[f"D{current_r}"] = proposal_data["pricing"]["currency"]
    ws[f"D{current_r}"].font = font_total_label
    ws[f"D{current_r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws[f"D{current_r}"].fill = fill_total

    ws[f"E{current_r}"] = proposal_data["pricing"]["client_total"]
    ws[f"E{current_r}"].font = font_total_val
    ws[f"E{current_r}"].alignment = Alignment(horizontal="right", vertical="center")
    ws[f"E{current_r}"].number_format = "$#,##0.00"
    ws[f"E{current_r}"].fill = fill_total

    for col in cols:
        ws[f"{col}{current_r}"].border = double_bottom

    ws.row_dimensions[current_r].height = 26
    current_r += 2

    # 5. Terms & Conditions Section
    ws.merge_cells(f"A{current_r}:E{current_r}")
    ws[f"A{current_r}"] = "TERMS & CONDITIONS OF CARRIAGE"
    ws[f"A{current_r}"].font = font_section_hdr
    current_r += 1

    for t in proposal_data["terms"]:
        ws.merge_cells(f"A{current_r}:E{current_r}")
        ws[f"A{current_r}"] = t
        ws[f"A{current_r}"].font = font_terms
        ws.row_dimensions[current_r].height = 15
        current_r += 1

    # Footer note
    current_r += 1
    ws.merge_cells(f"A{current_r}:E{current_r}")
    ws[f"A{current_r}"] = "RateSift Canada • PIPEDA Compliant • All rates calculated deterministically from verified tariffs."
    ws[f"A{current_r}"].font = Font(name="Arial", size=8, color="94A3B8")
    ws[f"A{current_r}"].alignment = Alignment(horizontal="center", vertical="center")

    # Set Column Widths
    col_widths = {
        "A": 16,
        "B": 24,
        "C": 36,
        "D": 14,
        "E": 18
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
