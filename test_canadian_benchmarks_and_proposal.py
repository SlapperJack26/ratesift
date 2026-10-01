import io
import openpyxl
from starlette.testclient import TestClient
from main import app
from services.ratesift_engine import quote_all_confirmed_carriers
from services.seed_canadian_benchmarks import seed_canadian_benchmarks_for_user
from services.auth_service import register_and_create_session

client = TestClient(app)

def test_canadian_benchmarks_and_proposal():
    print("\n" + "=" * 70)
    print("RUNNING CANADIAN BENCHMARKS & 1-CLICK CLIENT PROPOSAL TEST SUITE")
    print("=" * 70)

    # -------------------------------------------------------------
    # 1. VERIFY SEEDING OF CANADIAN BENCHMARK TARIFFS
    # -------------------------------------------------------------
    print("\n[STEP 1] Testing Canadian Benchmark Tariffs Seeding...")
    seeded_ids = seed_canadian_benchmarks_for_user("usr_alex_rivers", force=True)
    assert len(seeded_ids) == 4, f"Expected 4 benchmark tariffs, got {len(seeded_ids)}"
    print(f"   [OK] Seeded 4 active Canadian benchmark tariffs for Alex Rivers: {seeded_ids}")

    # -------------------------------------------------------------
    # 2. VERIFY MULTI-CARRIER SIDE-BY-SIDE QUOTING (CALGARY -> LINDSAY)
    # -------------------------------------------------------------
    print("\n[STEP 2] Testing Deterministic Multi-Carrier Quoting (Calgary -> Lindsay)...")
    res_calgary = quote_all_confirmed_carriers(
        user_id="usr_alex_rivers",
        origin="CALGARY, AB",
        destination="LINDSAY, ON",
        actual_weight=1450.0,
        accessorials=["liftgate"]
    )
    quotes = res_calgary["quotes"]
    assert len(quotes) >= 4, f"Expected at least 4 quotes, got {len(quotes)}"
    
    carrier_names = {q["carrier_name"] for q in quotes}
    assert "Day & Ross" in carrier_names
    assert "Manitoulin Transport" in carrier_names
    assert "Midland Transport" in carrier_names
    assert "Bison Transport" in carrier_names

    # Check pricing order (Rank 1 lowest price)
    for i in range(len(quotes) - 1):
        assert quotes[i]["final_total"] <= quotes[i + 1]["final_total"], "Quotes must be sorted by lowest price"

    top_carrier = quotes[0]
    second_carrier = quotes[1]
    print(f"   [OK] [1st Place] {top_carrier['carrier_name']} - ${top_carrier['final_total']:.2f} CAD ({top_carrier['transit_days']} days)")
    print(f"   [OK] [2nd Place] {second_carrier['carrier_name']} - ${second_carrier['final_total']:.2f} CAD ({second_carrier['transit_days']} days)")
    print(f"   [OK] Verified multi-carrier side-by-side ranking with {len(quotes)} carrier options")

    # -------------------------------------------------------------
    # 3. VERIFY TORONTO -> MONTREAL LANE (RULE 12 MINIMUM FLOORS)
    # -------------------------------------------------------------
    print("\n[STEP 3] Testing Corridors & Rule 12 Minimum Charge Floors (Toronto -> Montreal)...")
    res_toronto = quote_all_confirmed_carriers(
        user_id="usr_alex_rivers",
        origin="TORONTO, ON",
        destination="MONTREAL, QC",
        actual_weight=500.0,
        accessorials=["appointment"]
    )
    t_quotes = res_toronto["quotes"]
    assert len(t_quotes) >= 4
    for q in t_quotes:
        if q["carrier_name"] in ["Day & Ross", "Manitoulin Transport", "Midland Transport", "Bison Transport"]:
            assert q["carrier_min_charge"] > 0
            assert q["final_total"] >= q["carrier_min_charge"]
    print(f"   [OK] Successfully validated minimum charge comparisons on Toronto -> Montreal lane")

    # -------------------------------------------------------------
    # 4. VERIFY 1-CLICK CLIENT PROPOSAL PREVIEW ENDPOINT
    # -------------------------------------------------------------
    print("\n[STEP 4] Testing Client Proposal Preview API Endpoint...")
    sample_quote = quotes[0]
    payload = {
        "quote_id": res_calgary["quote_id"],
        "carrier_name": sample_quote["carrier_name"],
        "service_name": sample_quote["service_name"],
        "origin": sample_quote["origin"],
        "destination": sample_quote["destination"],
        "weight_lbs": sample_quote["billable_weight"],
        "base_rate": sample_quote["base_rate"],
        "total_surcharges": sample_quote["total_surcharges"],
        "surcharges": sample_quote["surcharges"],
        "transit_days": sample_quote["transit_days"],
        "currency": "CAD",
        "client_name": "Northern Manufacturing Ltd.",
        "markup_pct": 15.0
    }

    res = client.post(
        "/api/quotes/client-proposal/preview",
        json=payload,
        cookies={"ratesift_session": "session_demo_alex_rivers"}
    )
    assert res.status_code == 200, f"Preview failed: {res.text}"
    body = res.json()
    assert body["status"] == "success"
    prop = body["proposal"]
    assert prop["client"]["name"] == "Northern Manufacturing Ltd."
    assert prop["pricing"]["currency"] == "CAD"
    assert prop["pricing"]["effective_markup_pct"] == 15.0
    assert prop["pricing"]["client_total"] > sample_quote["final_total"]
    assert len(prop["line_items"]) >= 2  # Base transportation + liftgate
    print(f"   [OK] Proposal Preview generated: Client Total = ${prop['pricing']['client_total']:.2f} CAD (Margin: +${prop['pricing']['broker_margin_dollars']:.2f})")

    # -------------------------------------------------------------
    # 5. VERIFY 1-CLICK CLIENT PROPOSAL EXCEL EXPORT ENDPOINT
    # -------------------------------------------------------------
    print("\n[STEP 5] Testing Client Proposal Excel (.xlsx) Workbook Generation...")
    res_excel = client.post(
        "/api/quotes/client-proposal/excel",
        json=payload,
        cookies={"ratesift_session": "session_demo_alex_rivers"}
    )
    assert res_excel.status_code == 200, f"Excel generation failed: {res_excel.text}"
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res_excel.headers["content-type"]
    assert "RateSift_Client_Proposal_" in res_excel.headers["content-disposition"]

    # Validate workbook structure with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(res_excel.content))
    assert "Client_Freight_Proposal" in wb.sheetnames
    ws = wb["Client_Freight_Proposal"]
    
    # Check title banner and client name
    assert "TECHCORP LOGISTICS" in str(ws["A1"].value)
    assert "FREIGHT RATE QUOTE PROPOSAL" in str(ws["A2"].value)
    assert ws["B6"].value == "Northern Manufacturing Ltd."
    assert "CALGARY, AB" in str(ws["B11"].value)
    assert "LINDSAY, ON" in str(ws["B12"].value)
    print(f"   [OK] Downloaded and verified openpyxl Excel workbook: {len(res_excel.content)} bytes")

    # -------------------------------------------------------------
    # 6. VERIFY NEW USER AUTO-SEEDING UPON REGISTRATION
    # -------------------------------------------------------------
    print("\n[STEP 6] Testing New Broker User Registration Auto-Seeding...")
    new_user, token = register_and_create_session(
        name="Sarah Miller",
        email=f"sarah.miller.{uuid_hex()}@ontariofreight.ca",
        company="Ontario Freight Brokers",
        origin_zip="TORONTO, ON"
    )
    assert new_user["tier"] == "FREE"

    # Query quotes for the newly registered user
    res_new_user = quote_all_confirmed_carriers(
        user_id=new_user["id"],
        origin="TORONTO, ON",
        destination="OTTAWA, ON",
        actual_weight=850.0
    )
    assert len(res_new_user["quotes"]) >= 4, f"New user should have preloaded Canadian tariffs ready, got {len(res_new_user['quotes'])}"
    print(f"   [OK] New broker user '{new_user['name']}' instantly quoted {len(res_new_user['quotes'])} Canadian carriers without manual uploads!")

    print("\n" + "=" * 70)
    print("ALL CANADIAN BENCHMARK AND PROPOSAL GENERATOR TESTS PASSED (100%)")
    print("=" * 70)

def uuid_hex():
    import uuid
    return uuid.uuid4().hex[:6]

if __name__ == "__main__":
    test_canadian_benchmarks_and_proposal()
